import csv
import re
from pathlib import Path

import openpyxl

CSV_PATH = Path(r"c:\Users\User\Downloads\Заявки на билеты Каменное сердце - Libertad - Sheet1.csv")
XLSX_PATH = Path(r"c:\Users\User\Downloads\tickets_report_69eb348d56d27c0da363878e_20260610_165931.xlsx")


def norm_email(value):
    if value is None:
        return ""
    return str(value).strip().lower()


def norm_phone(value):
    if value is None:
        return ""
    digits = re.sub(r"\D", "", str(value))
    if len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    return digits


def load_waitlist():
    records = []
    seen_emails = set()
    with CSV_PATH.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            email = norm_email(row.get("Email"))
            if not email or email in seen_emails:
                continue
            seen_emails.add(email)
            records.append(
                {
                    "name": (row.get("Имя") or "").strip(),
                    "email": email,
                    "phone": norm_phone(row.get("Телефон")),
                }
            )
    return records


def load_tickets():
    workbook = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    header_idx = None
    headers = []
    for index, row in enumerate(rows):
        values = [str(cell).strip() if cell is not None else "" for cell in row]
        if any(value.lower() == "email" for value in values):
            header_idx = index
            headers = values
            break

    if header_idx is None:
        raise RuntimeError("Email column not found in ticket report")

    columns = {header: idx for idx, header in enumerate(headers) if header}
    email_idx = columns["Email"]
    phone_idx = next(
        (idx for header, idx in columns.items() if "тел" in header.lower()),
        None,
    )
    name_idx = next(
        (
            idx
            for header, idx in columns.items()
            if header in {"Покупатель", "Имя"} or "покуп" in header.lower()
        ),
        None,
    )

    tickets = []
    for row in rows[header_idx + 1 :]:
        if not any(row):
            continue
        email = norm_email(row[email_idx] if email_idx < len(row) else "")
        phone = norm_phone(row[phone_idx] if phone_idx is not None and phone_idx < len(row) else "")
        name = ""
        if name_idx is not None and name_idx < len(row) and row[name_idx]:
            name = str(row[name_idx]).strip()
        if not email and not phone:
            continue
        tickets.append({"name": name, "email": email, "phone": phone})

    return tickets


def main():
    waitlist = load_waitlist()
    tickets = load_tickets()

    ticket_emails = {ticket["email"] for ticket in tickets if ticket["email"]}
    ticket_phones = {ticket["phone"] for ticket in tickets if ticket["phone"]}

    matched = []
    for person in waitlist:
        by_email = person["email"] in ticket_emails
        by_phone = person["phone"] and person["phone"] in ticket_phones
        if by_email or by_phone:
            matched.append(
                {
                    **person,
                    "match": "email" if by_email else "phone",
                }
            )

    print(f"WAITLIST_UNIQUE={len(waitlist)}")
    print(f"TICKET_ROWS={len(tickets)}")
    print(f"MATCHED={len(matched)}")
    print("MATCHED_LIST_START")
    for person in matched:
        print(f"{person['name']}|{person['email']}|{person['phone']}|{person['match']}")
    print("MATCHED_LIST_END")

    unmatched = [person for person in waitlist if person["email"] not in {m["email"] for m in matched}]
    print(f"UNMATCHED={len(unmatched)}")
    print("UNMATCHED_LIST_START")
    for person in unmatched:
        print(f"{person['name']}|{person['email']}|{person['phone']}")
    print("UNMATCHED_LIST_END")


if __name__ == "__main__":
    main()
